#!/usr/bin/env python3
"""Reconcile the committed 2011 district PCA mirror against Census of India PCA SD."""
from __future__ import annotations

import json
import re
import sys
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "india-district-census-2011.json"
REPORT = ROOT / "reports" / "census-pca-2011-reconciliation.json"
URL = "https://censusindia.gov.in/nada/index.php/catalog/6191/download/9268/DDW_PCA0000_2011_Indiastatedist.xlsx"


def norm(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


STATE_ALIASES = {
    "orissa": "odisha", "uttaranchal": "uttarakhand",
    "nctofdelhi": "delhi", "nationalcapitalterritoryofdelhi": "delhi",
    "pondicherry": "puducherry", "jammukashmir": "jammuandkashmir",
    "jammuandkashmir": "jammuandkashmir",
    "damandiu": "damananddiu",
    "dadranagarhaveli": "dadraandnagarhaveli",
    "andamannicobarislands": "andamanandnicobarislands",
}
DISTRICT_ALIASES = {
    "leh": "lehladakh", "lehladakh": "lehladakh",
    "lahulspiti": "lahulandspiti", "lahulandspiti": "lahulandspiti",
    "mewat": "nuh", "gurgaon": "gurugram", "bangalore": "bengaluru",
    "bangalorerural": "bengalururural", "bellary": "ballari",
    "shimoga": "shivamogga", "mysore": "mysuru", "gulbarga": "kalaburagi",
    "pondicherry": "puducherry", "nicobars": "nicobar",
    "northmiddleandaman": "northandmiddleandaman",
}


def canon_state(value: object) -> str:
    key = norm(value)
    return STATE_ALIASES.get(key, key)


def canon_district(value: object) -> str:
    key = norm(value)
    return DISTRICT_ALIASES.get(key, key)


def number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return int(round(value))
    try:
        return int(round(float(str(value).replace(",", "").strip())))
    except (TypeError, ValueError):
        return None


def main():
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    records = json.loads(DATA.read_text(encoding="utf-8")).get("records", [])
    request = urllib.request.Request(URL, headers={"User-Agent": "ZipPrep-source-audit/1.0"})
    with urllib.request.urlopen(request, timeout=90) as response:
        workbook_bytes = response.read()
    tmp = ROOT / "reports" / "official-pca-sd-2011.xlsx"
    tmp.write_bytes(workbook_bytes)
    workbook = load_workbook(tmp, read_only=True, data_only=True)
    sheet = workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.worksheets[0]
    iterator = sheet.iter_rows(values_only=True)
    header = next(iterator, None)
    if not header:
        raise RuntimeError("Official PCA workbook has no header row")
    columns = {str(value).strip(): index for index, value in enumerate(header) if value is not None}
    required = {"State", "District", "Level", "Name", "TRU", "No_HH", "TOT_P", "TOT_M", "TOT_F", "P_06", "P_SC", "P_ST", "P_LIT", "TOT_WORK_P"}
    missing_columns = sorted(required - set(columns))
    if missing_columns:
        report = {
            "status": "blocked_schema_unrecognised",
            "sourceUrl": URL,
            "sourceCatalogue": "https://censusindia.gov.in/nada/index.php/catalog/6191",
            "checkedAt": datetime.now(timezone.utc).isoformat(),
            "workbookHeaders": list(columns),
            "missingRequiredColumns": missing_columns,
            "mirrorRows": len(records),
            "message": "The official PCA workbook schema changed; no records were verified."
        }
        REPORT.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
        print(json.dumps(report, indent=2))
        return 2

    rows = list(iterator)
    state_names = {}
    for row in rows:
        get = lambda field: row[columns[field]] if columns[field] < len(row) else None
        if str(get("Level") or "").strip().upper() == "STATE" and str(get("TRU") or "").strip().lower() == "total":
            state_names[str(get("State")).zfill(2)] = str(get("Name")).strip()

    official, official_by_code, duplicate_official = {}, {}, []
    for row in rows:
        get = lambda field: row[columns[field]] if columns[field] < len(row) else None
        if str(get("Level") or "").strip().upper() != "DISTRICT" or str(get("TRU") or "").strip().lower() != "total":
            continue
        state_name = state_names.get(str(get("State")).zfill(2))
        district_name = str(get("Name") or "").strip()
        if not state_name or not district_name:
            continue
        key = (canon_state(state_name), canon_district(district_name))
        pop, male, female = number(get("TOT_P")), number(get("TOT_M")), number(get("TOT_F"))
        literate, children_0_6 = number(get("P_LIT")), number(get("P_06"))
        values = {
            "households": number(get("No_HH")),
            "population": pop, "male": male, "female": female,
            "literate": literate, "sc": number(get("P_SC")),
            "st": number(get("P_ST")), "workers": number(get("TOT_WORK_P")),
            "literacyRate": round(literate / (pop - children_0_6) * 100, 2) if pop is not None and children_0_6 is not None and pop > children_0_6 and literate is not None else None,
            "sexRatio": round(female / male * 1000) if male else None,
            "censusStateCode": str(get("State")).zfill(2),
            "censusDistrictCode": str(get("District")).zfill(3),
        }
        if key in official:
            duplicate_official.append({"state": key[0], "district": key[1]})
        official[key] = values
        official_by_code[values["censusDistrictCode"]] = values

    mismatches, missing, matched = [], [], 0
    fields = ("households", "population", "male", "female", "literate", "sc", "st", "workers")
    for record in records:
        key = (canon_state(record.get("state")), canon_district(record.get("district")))
        source = official.get(key)
        if source is None and record.get("censusDistrictCode") is not None:
            source = official_by_code.get(str(record.get("censusDistrictCode")).zfill(3))
        if source is None:
            missing.append({"state": record.get("state"), "district": record.get("district"), "censusDistrictCode": record.get("censusDistrictCode")})
            continue
        differences = []
        for field in fields:
            if source.get(field) is not None and number(record.get(field)) != source[field]:
                differences.append({"field": field, "mirror": number(record.get(field)), "official": source[field]})
        for field in ("literacyRate", "sexRatio"):
            value = source.get(field)
            if value is not None:
                try:
                    if abs(float(record.get(field)) - float(value)) > (0.02 if field == "literacyRate" else 0.5):
                        differences.append({"field": field, "mirror": record.get(field), "official": value})
                except (TypeError, ValueError):
                    differences.append({"field": field, "mirror": record.get(field), "official": value})
        if record.get("censusDistrictCode") is not None and str(record.get("censusDistrictCode")).zfill(3) != source["censusDistrictCode"]:
            differences.append({"field": "censusDistrictCode", "mirror": record.get("censusDistrictCode"), "official": source["censusDistrictCode"]})
        matched += 1
        if differences:
            mismatches.append({"state": record.get("state"), "district": record.get("district"), "censusDistrictCode": record.get("censusDistrictCode"), "differences": differences})

    # Generate a source-normalised copy for review and integration.
    reconciled = json.loads(json.dumps(json.loads(DATA.read_text(encoding="utf-8"))))
    for record in reconciled.get("records", []):
        code = str(record.get("censusDistrictCode", "")).zfill(3)
        source = official_by_code.get(code)
        if source:
            for field in ("households", "population", "male", "female", "literate", "literacyRate", "sexRatio", "sc", "st", "workers"):
                if source.get(field) is not None:
                    record[field] = source[field]
            record["source"] = "Census of India, Primary Census Abstract (PCA), 2011"
            record["sourceUrl"] = "https://censusindia.gov.in/nada/index.php/catalog/6191"
            record["sourceDownloadUrl"] = URL
            record["sourceStatus"] = "official_source_reconciled"
    reconciled["source"] = "Census of India, Primary Census Abstract (PCA), 2011"
    reconciled["sourceUrl"] = "https://censusindia.gov.in/nada/index.php/catalog/6191"
    reconciled["sourceDownloadUrl"] = URL
    reconciled["reconciliationPolicy"] = "Historical district totals matched to official 2011 PCA by original Census district code; no current-boundary extrapolation."
    canonical_path = ROOT / "reports" / "india-district-census-2011-official-reconciled.json"
    canonical_path.write_text(json.dumps(reconciled, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    status = "source_code_reconciled" if len(records) == 640 and matched == 640 and not missing and not duplicate_official else "review_required"
    report = {
        "status": status, "sourceUrl": URL,
        "sourceCatalogue": "https://censusindia.gov.in/nada/index.php/catalog/6191",
        "checkedAt": datetime.now(timezone.utc).isoformat(), "workbookSheet": sheet.title,
        "workbookDistrictTotalRows": len(official), "mirrorRows": len(records), "matchedRows": matched,
        "missingRows": missing, "mismatchRows": mismatches, "duplicateOfficialKeys": duplicate_official,
        "reconciledDataset": "reports/india-district-census-2011-official-reconciled.json",
        "matchedFields": list(fields) + ["literacyRate", "sexRatio", "censusDistrictCode"],
        "policy": "Technical comparison only. Application data remains unverified until discrepancies and historical boundary/name changes are reviewed and the reviewed result is committed."
    }
    REPORT.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({
        "status": status, "workbookDistrictTotalRows": len(official),
        "mirrorRows": len(records), "matchedRows": matched,
        "missingRows": len(missing), "mismatchRows": len(mismatches),
        "duplicateOfficialKeys": len(duplicate_official)
    }, indent=2))
    return 0 if status == "source_code_reconciled" else 1


if __name__ == "__main__":
    sys.exit(main())
